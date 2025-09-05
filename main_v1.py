import os
import pandas as pd
from typing import Dict, List
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from medicare_v1 import scrape_medicare as scrape_medicare_v1
from yellowpages_v3 import scrape_yellowpages
from gmaps_v1 import scrape_google_maps

# from browser import driver
from config import _outfile_from_location, results_dir_for_today

def format_worksheet(worksheet, num_columns, num_rows):
    orange_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    header_font = Font(name="Arial", size=14, bold=True)
    data_font = Font(name="Arial", size=12)
    
    for col in range(1, num_columns + 1):
        worksheet.column_dimensions[get_column_letter(col)].width = 40
        
        header_cell = worksheet.cell(row=1, column=col)
        header_cell.fill = orange_fill
        header_cell.font = header_font
        
        for row in range(2, num_rows + 2):  
            data_cell = worksheet.cell(row=row, column=col)
            data_cell.font = data_font

def run(
    location: str,
    yp_keyword: str,
    gmaps_keyword: str,
    limit: int,
    radius_miles: int,
    sources: Dict[str, bool],
    medicare_categories: Dict[str, bool],
) -> Dict:
    """
    Execute selected scrapers and write Excel.
    Returns a summary dict for the UI.
    """
    cols = ["Name", "Mileage", "Address", "City", "State", "Zip", "Phone Number", "Specialty", "Full Address"]
    rows_medicare: List = []
    rows_yp: List = []
    rows_gmaps: List = []
    
    output_name = _outfile_from_location(location)         # from config.py
    results_dir = results_dir_for_today()                  # from config.py
    output_file = os.path.join(results_dir, output_name)
    try:
        if any(medicare_categories.values()):
            rows_medicare = scrape_medicare_v1(location, medicare_categories, limit, radius_miles)
            
        if sources.get("YellowPages"):
            rows_yp = scrape_yellowpages(location, yp_keyword, limit, radius_miles)
            
        if sources.get("GoogleMaps"):
            rows_gmaps = scrape_google_maps(location, gmaps_keyword, limit, radius_miles)

        # Build DataFrames
        df_medicare = pd.DataFrame(rows_medicare, columns=cols) if rows_medicare else pd.DataFrame(columns=cols)
        df_yp = pd.DataFrame(rows_yp, columns=cols) if rows_yp else pd.DataFrame(columns=cols)
        df_gmaps = pd.DataFrame(rows_gmaps, columns=cols) if rows_gmaps else pd.DataFrame(columns=cols)

        # Remove old file if present
        if os.path.exists(output_file):
            os.remove(output_file)

        # Write only non-empty sheets
        with pd.ExcelWriter(output_file, engine="openpyxl") as xw:
            if not df_medicare.empty:
                df_medicare.to_excel(xw, sheet_name="Medicare", index=False)
                format_worksheet(xw.sheets["Medicare"], len(df_medicare.columns), len(df_medicare))
            
            if not df_yp.empty:
                df_yp.to_excel(xw, sheet_name="YellowPages", index=False)
                format_worksheet(xw.sheets["YellowPages"], len(df_yp.columns), len(df_yp))
            
            if not df_gmaps.empty:
                df_gmaps.to_excel(xw, sheet_name="GoogleMaps", index=False)
                format_worksheet(xw.sheets["GoogleMaps"], len(df_gmaps.columns), len(df_gmaps))

        summary = {
            "outfile": output_file,
            "counts": {
                "Medicare": len(df_medicare),
                "YellowPages": len(df_yp),
                "GoogleMaps": len(df_gmaps),
            },
        }
        print(f"\nSaved to: {output_file}")
        print(f"Summary: {summary}")
        for k, v in summary["counts"].items():
            print(f"  {k:11s}: {v} rows")
        return summary
    except Exception as e:
        print(f"Error occurred: {e}")
        print(f"Error type: {type(e)}")
        import traceback
        traceback.print_exc()
        return None
